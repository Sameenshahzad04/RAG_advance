import os
import logging
from typing import List, Dict, Any

logger = logging.getLogger(__name__)


def extract_content(file_path: str) -> List[Dict[str, Any]]:
    """
    Main entry point — detects file type and calls the right extractor.
    Supported: .pdf, .docx, .txt, .md
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext in (".txt", ".md"):
        return _extract_txt(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")


def _extract_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extract text and tables from a PDF file using pdfplumber (or pypdf fallback)."""
    elements: List[Dict[str, Any]] = []

    # Try pdfplumber first (better table extraction)
    try:
        import pdfplumber  # type: ignore

        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                # Extract tables from this page
                tables = page.extract_tables()
                if tables:
                    for tbl in tables:
                        md = _format_table_markdown(tbl)
                        if md:
                            elements.append({
                                "type": "table",
                                "content": f"[Table - Page {page_num}]\n{md}"
                            })

                # Extract text from this page
                # FIX: x_tolerance and y_tolerance force pdfplumber to detect gaps between words
                text = page.extract_text(x_tolerance=2, y_tolerance=3)
                if text and text.strip():
                    elements.append({
                        "type": "text",
                        "content": f"[Page {page_num}]\n{text.strip()}"
                    })

        if elements:
            return elements

    except Exception as e:
        logger.warning(f"pdfplumber failed ({e}), trying pypdf...")

    # Fallback to pypdf (simpler, text-only extraction)
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(file_path)
        for idx, page in enumerate(reader.pages, start=1):
           
           # FIX: extraction_mode="layout" preserves word spaces on LaTeX/Beamer PDFs
            try:
                text = page.extract_text(extraction_mode="layout")
            except Exception:
                text = page.extract_text()
            if text and text.strip():
                elements.append({
                    "type": "text",
                    "content": f"[Page {idx}]\n{text.strip()}"
                })
    except Exception as e:
        logger.error(f"pypdf also failed: {e}")
        raise RuntimeError(f"Cannot extract PDF: {e}")

    return elements


def _extract_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extract paragraphs and tables from a DOCX file."""
    import docx  # type: ignore

    doc = docx.Document(file_path)
    elements: List[Dict[str, Any]] = []

    # Extract all paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text and paragraph.text.strip():
            elements.append({"type": "text", "content": paragraph.text.strip()})

    # Extract all tables
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _format_table_markdown(rows)
        if md:
            elements.append({"type": "table", "content": f"[Table]\n{md}"})

    return elements


def _extract_txt(file_path: str) -> List[Dict[str, Any]]:
    """Read a plain text file and return its content."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    return [{"type": "text", "content": content.strip()}]


def _format_table_markdown(table_data: List[List[Any]]) -> str:
    """
    Convert a 2D list (rows x cols) into a Markdown table string.
    Example output:
        | Name | Age |
        | --- | --- |
        | Alice | 30 |
    """
    # Clean up: convert cells to strings, remove empty rows
    rows = [
        [str(cell or "").replace("\n", " ").strip() for cell in row]
        for row in table_data
        if any(row)
    ]
    if not rows:
        return ""

    # Make all rows the same width
    num_cols = max(len(row) for row in rows)
    for row in rows:
        while len(row) < num_cols:
            row.append("")

    # Build markdown table
    header = rows[0]
    separator = ["---"] * num_cols
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)







